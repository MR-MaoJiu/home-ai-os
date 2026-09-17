"""生成有效的文档格式测试文件；用真实解析器读取，不提供解析结果替身。"""
import sys
from pathlib import Path
from pptx import Presentation
from PIL import Image, ImageDraw, ImageFont

root=Path(sys.argv[1]);root.mkdir(parents=True,exist_ok=True)
text='Saturday backup at 3 PM'
stream=f'BT /F1 18 Tf 50 750 Td ({text}) Tj ET'.encode()
objects=[b'<< /Type /Catalog /Pages 2 0 R >>',b'<< /Type /Pages /Kids [3 0 R] /Count 1 >>',b'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>',b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>',b'<< /Length '+str(len(stream)).encode()+b' >>\nstream\n'+stream+b'\nendstream']
data=bytearray(b'%PDF-1.4\n');offsets=[0]
for i,obj in enumerate(objects,1):
    offsets.append(len(data));data.extend(f'{i} 0 obj\n'.encode()+obj+b'\nendobj\n')
xref=len(data);data.extend(f'xref\n0 {len(objects)+1}\n0000000000 65535 f \n'.encode())
for offset in offsets[1:]:data.extend(f'{offset:010d} 00000 n \n'.encode())
data.extend(f'trailer << /Size {len(objects)+1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n'.encode())
(root/'sample.pdf').write_bytes(data)
slides=Presentation();slide=slides.slides.add_slide(slides.slide_layouts[1]);slide.shapes.title.text='家庭会议记录';slide.placeholders[1].text='周六下午三点检查备份。';slides.save(root/'sample.pptx')
image=Image.new('RGB',(1200,300),'white');ImageDraw.Draw(image).text((40,100),text,font=ImageFont.load_default(size=52),fill='black');image.save(root/'sample.png')
